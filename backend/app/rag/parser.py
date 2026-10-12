import io
import re
from typing import List, Optional
from pydantic import BaseModel, Field
import pypdf
import docx


class DocumentSection(BaseModel):
    """Represents a discrete section or page from a parsed document."""
    text: str
    page_number: Optional[int] = None
    section_type: str = "text"  # "text", "table", "heading"
    metadata: dict = Field(default_factory=dict)


class ParsedDocument(BaseModel):
    """Standardized representation of a parsed multi-format document."""
    filename: str
    mime_type: str
    total_pages: int = 1
    sections: List[DocumentSection] = Field(default_factory=list)

    @property
    def full_text(self) -> str:
        """Convenience property joining all section texts with double newlines."""
        return "\n\n".join(section.text.strip() for section in self.sections if section.text.strip())


def normalize_text(text: str) -> str:
    """Clean and normalize whitespace, remove non-printable characters."""
    if not text:
        return ""
    # Remove control characters except newline and tab
    text = re.sub(r"[\x00-\x08\x0B-\x0C\x0E-\x1F\x7F]", "", text)
    # Replace non-breaking spaces with standard space
    text = text.replace("\xa0", " ")
    # Replace multiple spaces with a single space
    text = re.sub(r"[ \t]+", " ", text)
    # Collapse 3 or more consecutive newlines into 2
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_pdf(file_bytes: bytes, filename: str) -> ParsedDocument:
    """Parse PDF document and extract text per page with 1-based page numbering."""
    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    total_pages = len(reader.pages)
    sections: List[DocumentSection] = []

    for page_idx, page in enumerate(reader.pages):
        page_num = page_idx + 1
        page_text = page.extract_text() or ""
        cleaned = normalize_text(page_text)
        if cleaned:
            sections.append(
                DocumentSection(
                    text=cleaned,
                    page_number=page_num,
                    section_type="text",
                    metadata={"page": page_num, "total_pages": total_pages},
                )
            )

    return ParsedDocument(
        filename=filename,
        mime_type="application/pdf",
        total_pages=total_pages,
        sections=sections,
    )


def _table_to_markdown(table: docx.table.Table) -> str:
    """Convert a python-docx Table into a formatted Markdown table."""
    rows_data = []
    for row in table.rows:
        row_text = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        # Only add row if it has at least some content
        if any(row_text):
            rows_data.append(row_text)

    if not rows_data:
        return ""

    # Ensure all rows have equal column length
    max_cols = max(len(r) for r in rows_data)
    padded_rows = [r + [""] * (max_cols - len(r)) for r in rows_data]

    md_lines = []
    # Header row
    header = padded_rows[0]
    md_lines.append("| " + " | ".join(header) + " |")
    md_lines.append("| " + " | ".join(["---"] * max_cols) + " |")

    # Data rows
    for row in padded_rows[1:]:
        md_lines.append("| " + " | ".join(row) + " |")

    return "\n".join(md_lines)


def parse_docx(file_bytes: bytes, filename: str) -> ParsedDocument:
    """Parse DOCX document, extracting paragraphs and converting tables to Markdown."""
    doc = docx.Document(io.BytesIO(file_bytes))
    sections: List[DocumentSection] = []

    # Iterate through body elements preserving document order
    current_text_buffer: List[str] = []

    def flush_text_buffer():
        if current_text_buffer:
            combined = normalize_text("\n\n".join(current_text_buffer))
            if combined:
                sections.append(
                    DocumentSection(
                        text=combined,
                        page_number=1,
                        section_type="text",
                    )
                )
            current_text_buffer.clear()

    for element in doc.element.body:
        # Paragraph element
        if element.tag.endswith("p"):
            p = docx.text.paragraph.Paragraph(element, doc)
            text = p.text.strip()
            if text:
                current_text_buffer.append(text)
        # Table element
        elif element.tag.endswith("tbl"):
            flush_text_buffer()
            table = docx.table.Table(element, doc)
            table_md = _table_to_markdown(table)
            if table_md:
                sections.append(
                    DocumentSection(
                        text=table_md,
                        page_number=1,
                        section_type="table",
                        metadata={"type": "markdown_table"},
                    )
                )

    flush_text_buffer()

    return ParsedDocument(
        filename=filename,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        total_pages=1,
        sections=sections,
    )


def parse_text(file_bytes: bytes, filename: str, mime_type: str = "text/plain") -> ParsedDocument:
    """Parse plain text or Markdown document using UTF-8 decoding."""
    # Attempt UTF-8 with fallback
    try:
        raw_text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        raw_text = file_bytes.decode("utf-8", errors="replace")

    cleaned = normalize_text(raw_text)
    sections: List[DocumentSection] = []

    if cleaned:
        # Split by top-level Markdown headers if it's markdown, otherwise single or paragraph section
        is_markdown = filename.lower().endswith(".md") or "markdown" in mime_type
        if is_markdown and re.search(r"^#{1,3}\s+", cleaned, re.MULTILINE):
            # Split sections by markdown headings
            parts = re.split(r"(?m)(?=^#{1,3}\s+)", cleaned)
            for part in parts:
                part_clean = part.strip()
                if part_clean:
                    sections.append(
                        DocumentSection(
                            text=part_clean,
                            page_number=1,
                            section_type="heading" if part_clean.startswith("#") else "text",
                        )
                    )
        else:
            sections.append(
                DocumentSection(
                    text=cleaned,
                    page_number=1,
                    section_type="text",
                )
            )

    return ParsedDocument(
        filename=filename,
        mime_type=mime_type,
        total_pages=1,
        sections=sections,
    )


def parse_document(
    file_bytes: bytes,
    filename: str,
    mime_type: Optional[str] = None,
) -> ParsedDocument:
    """
    Main entry point for multi-format document parsing.
    Supports PDF, DOCX, TXT, and Markdown.
    """
    if not file_bytes:
        raise ValueError("Cannot parse empty file content")

    ext = filename.lower().split(".")[-1] if "." in filename else ""
    detected_mime = (mime_type or "").lower()

    if ext == "pdf" or "pdf" in detected_mime:
        parsed = parse_pdf(file_bytes, filename)
    elif ext == "docx" or "wordprocessingml" in detected_mime:
        parsed = parse_docx(file_bytes, filename)
    elif ext in ("txt", "md") or "text" in detected_mime:
        parsed = parse_text(file_bytes, filename, mime_type=mime_type or "text/plain")
    else:
        # Try text parser as default fallback
        try:
            parsed = parse_text(file_bytes, filename, mime_type=mime_type or "text/plain")
        except Exception as e:
            raise ValueError(f"Unsupported file format '{ext}' for parsing: {e}") from e

    if not parsed.sections or not parsed.full_text.strip():
        raise ValueError(f"No extractable text found in '{filename}'. Ensure document is not blank or scanned image.")

    return parsed
