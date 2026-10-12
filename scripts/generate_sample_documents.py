"""
Generate diverse real-world sample documents (DOCX, PDF) in data/sample_documents/
for RAG ingestion and retrieval testing.
"""

import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


OUTPUT_DIR = Path(__file__).resolve().parent.parent / "data" / "sample_documents"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def create_docx_sample():
    docx_path = OUTPUT_DIR / "chinh_sach_nhan_su_va_cong_tac_phi.docx"
    doc = docx.Document()

    # Title
    title = doc.add_heading("CHINH SACH NHAN SU VA DINH MUC CONG TAC PHI 2026", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_meta = doc.add_paragraph("Ma van ban: CS-NS-2026/05 | Ap dung: Toan the can bo cong nhan vien ERP")
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("1. Nguyen tac chung ve cong tac phi", level=1)
    doc.add_paragraph(
        "Cong tac phi la khoan chi phi ho tro can bo nhan vien khi duoc cu di cong tac ben ngoai "
        "dia ban tinh/thanh pho noi lam viec thuong xuyen. Moi chuyen cong tac phai co Ke hoach cong tac "
        "duoc phe duyet truoc it nhat 03 ngay lam viec tren he thong ERP."
    )

    doc.add_heading("2. Bang dinh muc chi phi luu tru va phu cap", level=1)
    doc.add_paragraph("Dinh muc chi tieu duoc ap dung theo cap bac nhan su nhu sau:")

    # Table with 4 columns
    table = doc.add_table(rows=4, cols=4)
    table.style = "Table Grid"

    headers = ["Chuc danh", "Tien phong / dem (VND)", "Phu cap luu tru / ngay", "Phuong tien di chuyen"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True

    data = [
        ["Nhan vien / Sales Rep", "800,000", "300,000", "Tau hoa / Xe khach / May bay > 500km"],
        ["Truong phong / Quan ly", "1,500,000", "600,000", "May bay pho thong (Economy)"],
        ["Ban Giam doc (C-Level)", "3,000,000", "1,200,000", "May bay thuong gia / Pho thong dac biet"],
    ]

    for row_idx, row_data in enumerate(data, start=1):
        for col_idx, text in enumerate(row_data):
            table.cell(row_idx, col_idx).text = text

    doc.add_heading("3. Ho so va thoi han thanh quyet toan", level=1)
    doc.add_paragraph(
        "Trong vong 07 ngay lam viec ke tu ngay ket thuc dot cong tac, nhan su phai hoan tat ho so "
        "thanh toan gom: Giay di duong co xac nhan cua khach hang hoac chi nhanh, Hoa don do VAT ve may bay/khach san. "
        "Qua thoi han 07 ngay tren ma khong co ly do chinh dang, phong Ke toan se tu choi chi tra."
    )

    doc.save(docx_path)
    print(f"[+] Created DOCX sample: {docx_path}")


def create_pdf_sample():
    pdf_path = OUTPUT_DIR / "so_tay_an_toan_lao_dong_va_kho_hang.pdf"
    doc = SimpleDocTemplate(str(pdf_path), pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()

    story = []

    # Custom styles
    title_style = ParagraphStyle("TitleStyle", parent=styles["Heading1"], fontSize=18, leading=22, alignment=1, textColor=colors.HexColor("#1A365D"))
    heading_style = ParagraphStyle("HeadingStyle", parent=styles["Heading2"], fontSize=14, leading=18, textColor=colors.HexColor("#2B6CB0"))
    body_style = ParagraphStyle("BodyStyle", parent=styles["Normal"], fontSize=10, leading=14)

    # --- PAGE 1 ---
    story.append(Paragraph("SO TAY QUY TRINH AN TOAN VAN HANH KHO TONG ERP 2026", title_style))
    story.append(Spacer(1, 15))
    story.append(Paragraph("Pham vi ap dung: Kho Tong Ha Noi (WH-HN) va Kho Tong Ho Chi Minh (WH-HCM)", body_style))
    story.append(Spacer(1, 15))

    story.append(Paragraph("Phan 1: Quy dinh trang bi bao ho lao dong ca nhan (PPE)", heading_style))
    story.append(Paragraph(
        "Moi nhan vien kho, lai xe nang va khach hang khi buoc vao khu vuc kho hang bat buoc phai trang bi day du: "
        "<br/>1. Giay bao ho mui thep chong dap ngon va chong dinh. "
        "<br/>2. Ao phan quang tieu chuan mau cam hoac xanh neon. "
        "<br/>3. Mu bao ho cung dat chuan TCVN tai cac khu vuc ke hang cao tang (Khu vuc Tier-High)."
        "<br/>Nghiem cam di dep le, giay the thao vai hoac khong mac ao phan quang vao kho.",
        body_style
    ))
    story.append(Spacer(1, 15))

    story.append(Paragraph("Phan 2: Toc do gioi han cua phuong tien nang ha (Forklift)", heading_style))
    story.append(Paragraph(
        "Toc do toi da cho phep cua xe nang Forklift di chuyen ben trong hanh lang kho chinh la khong qua <b>10 km/h</b>. "
        "Khi di qua cac nga ba hanh lang, cua ra vao hoac khu vuc tap ket hang hoa, toc do toi da la khong qua <b>5 km/h</b> "
        "va bat buoc phai bam coi canh bao truoc khi doi huong.",
        body_style
    ))

    # --- PAGE 2 ---
    story.append(PageBreak())

    story.append(Paragraph("Phan 3: Quy trinh ung pho su co va lien he khan cap", heading_style))
    story.append(Paragraph(
        "Khi phat hien su co chay no, chap dien hoac do vo hang hoa hoa chat, can thuc hien ngay 3 buoc sau: "
        "<br/>- <b>Buoc 1:</b> An chuong bao dong hoa hoan gan nhat va ngat cau dao dien tong tai tram phu trach. "
        "<br/>- <b>Buoc 2:</b> So tan toan bo nhan su theo loi thoat hiem da danh dau son mau vang tren san kho. "
        "<br/>- <b>Buoc 3:</b> Goi ngay duong day nong khan cap an toan kho ERP: <b>1900-6868 (Phim 2)</b>.",
        body_style
    ))
    story.append(Spacer(1, 15))

    story.append(Paragraph("Phan 4: Thoi han bao cao su co tren he thong ERP", heading_style))
    story.append(Paragraph(
        "Truong ca truc hoac Thu kho phai lap Bien ban ghi nhan su co va nhap thong tin len module "
        "Audit & Incident tren he thong ERP Copilot trong vong toi da <b>02 gio</b> ke tu khi su co duoc khac phuc tam thoi.",
        body_style
    ))

    doc.build(story)
    print(f"[+] Created PDF sample: {pdf_path}")


if __name__ == "__main__":
    create_docx_sample()
    create_pdf_sample()
